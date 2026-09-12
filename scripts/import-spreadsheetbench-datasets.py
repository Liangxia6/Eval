#!/usr/bin/env python3
"""Import 12 curated, pinned SpreadsheetBench Verified cases without editing XLSX bytes.

Requires Python 3 and openpyxl (read-only extraction). The official archive is
downloaded only with --download. Existing differing files are never overwritten.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
from pathlib import Path

from compact_dataset_bundles import SPREADSHEET_LICENSE, compact_bundle
import re
import tarfile
import urllib.request

import openpyxl
from openpyxl.utils.cell import range_boundaries

REPOSITORY = "https://github.com/RUCKBReasoning/SpreadsheetBench"
COMMIT = "49b73a94775fb489063f60ca1865e3a650079a79"
ARCHIVE_PATH = "data/spreadsheetbench_verified_400.tar.gz"
ARCHIVE_SHA256 = "10ef893dd29cb13ab97143ea787e68cdc9574a13873ab9a54e50b31dc03fc949"
ARCHIVE_ROOT = "spreadsheetbench_verified_400"
URL = f"https://raw.githubusercontent.com/RUCKBReasoning/SpreadsheetBench/{COMMIT}/{ARCHIVE_PATH}"
LABELS = ["tool-data", "tool-document", "reasoning-planning", "artifact-delivery"]
# Explicit IDs make the selection stable. Six sheet tasks and six cell tasks;
# avoid today's-date dependence, malformed ranges and VBA-only deliverables.
SELECTED = {
    "493-18": "按另一列的匹配结果清理数据并保留无关列",
    "66-24": "按相对最大日期的时间窗口筛选跨表记录",
    "82-38": "跨工作表匹配发票编号并按分段重复填充",
    "84-40": "跨工作表汇总两列金额并生成汇总表",
    "178-22": "根据多列 OR 条件提取记录到另一工作表",
    "463-17": "根据数量列展开并重复公司名称",
    "12307": "根据公司和员工数统计有业务的国家数量",
    "30930": "按标记分段累计正数并在指定行填写结果",
    "38823": "结合日期区间与文本包含条件求和",
    "40478": "按空格和连字符边界用公式抽取文本",
    "41601": "根据学生名称动态读取对应工作表中的年龄",
    "44389": "检索每行最小正数对应的列标题并处理并列",
}


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def write(path: Path, data: bytes) -> None:
    if path.exists() and path.read_bytes() != data:
        raise ValueError(f"refusing to overwrite changed file: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)


def json_bytes(value) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode()


def answer_text(expected) -> str:
    """Keep the full expected result in the template's open-text answer field."""
    return ("output/result.xlsx 应与 private/reference.xlsx 的目标区域结果一致。"
            "以下 JSON 列出全部目标单元格及参考信息；"
            "相同结构也保存在 private/expected.json，供评分侧读取。\n\n"
            + json.dumps(expected, ensure_ascii=False, indent=2, allow_nan=False))


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return {"type": type(value).__name__, "iso": value.isoformat()}
    if isinstance(value, dt.timedelta):
        return {"type": "timedelta", "seconds": value.total_seconds()}
    return value


def expected_regions(record, workbook_bytes):
    """Extract cached answers from original gold bytes; never recalculate gold."""
    values = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(workbook_bytes), data_only=False)
    regions = []
    for position in record["answer_position"].split(","):
        if "!" in position:
            sheet, address = position.rsplit("!", 1)
            sheet = sheet.strip("'")
        else:
            sheet = record.get("answer_sheet", values.sheetnames[0])
            address = position
        address = address.strip("'")
        if sheet not in values.sheetnames or not re.fullmatch(r"[A-Z]+[1-9][0-9]*(?::[A-Z]+[1-9][0-9]*)?", address):
            raise ValueError(f"ambiguous target {record['id']}: {sheet}!{address}")
        min_col, min_row, max_col, max_row = range_boundaries(address)
        if (max_col - min_col + 1) * (max_row - min_row + 1) > 5000:
            raise ValueError("target too large for curated batch")
        cells = {}
        for row in values[sheet].iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if cell.data_type == "e":
                    raise ValueError(f"gold contains error {record['id']}:{cell.coordinate}")
                original = formulas[sheet][cell.coordinate]
                # Excel stores a calculated empty-string formula as t="str";
                # openpyxl exposes that legitimate empty result as None.
                if original.data_type == "f" and cell.value is None and cell.data_type != "str":
                    raise ValueError(f"gold lacks cached formula result {record['id']}:{cell.coordinate}")
                cells[cell.coordinate] = json_value(cell.value)
        regions.append({"sheet": sheet, "range": address, "cells": cells})
    values.close()
    formulas.close()
    return regions


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", type=Path, required=True)
    parser.add_argument("--output-root", type=Path, default=Path(__file__).resolve().parents[1] / "datasets" / "spreadsheetbench")
    parser.add_argument("--download", action="store_true")
    args = parser.parse_args()
    if args.download and not args.archive.exists():
        args.archive.parent.mkdir(parents=True, exist_ok=True)
        with urllib.request.urlopen(URL, timeout=120) as response:
            write(args.archive, response.read())
    archive_bytes = args.archive.read_bytes()
    if digest(archive_bytes) != ARCHIVE_SHA256:
        raise ValueError("archive does not match pinned SHA-256")
    with tarfile.open(fileobj=io.BytesIO(archive_bytes), mode="r:gz") as archive:
        def member(relative):
            item = archive.getmember(f"{ARCHIVE_ROOT}/{relative}")
            if not item.isfile():
                raise ValueError(f"expected regular archive member: {relative}")
            return archive.extractfile(item).read()

        metadata_bytes = member("dataset.json")
        rows = {str(row["id"]): row for row in json.loads(metadata_bytes)}
        prepared = []
        for task_id, summary in SELECTED.items():
            record = rows[task_id]
            prefix = f"spreadsheet/{task_id}"
            init_path = f"{prefix}/1_{task_id}_init.xlsx"
            gold_path = f"{prefix}/1_{task_id}_golden.xlsx"
            prompt_path = f"{prefix}/prompt.txt"
            init, gold, prompt = member(init_path), member(gold_path), member(prompt_path)
            if prompt.decode().strip() != record["instruction"].strip():
                raise ValueError(f"prompt disagrees with metadata: {task_id}")
            # Read every input to catch corrupt workbooks while preserving bytes.
            workbook = openpyxl.load_workbook(io.BytesIO(init), read_only=True, data_only=False)
            sheets = workbook.sheetnames
            workbook.close()
            targets = expected_regions(record, gold)
            title = f"spreadsheetbench-{task_id}"
            question = {
                "schema": "dsheval.question/v1", "id": f"spreadsheetbench.{title}",
                "version": "1.0.0", "title": title,
                "matching": {"datasetId": f"dataset.spreadsheetbench.{title}/v1", "description": f"{summary}。输入为真实 Excel 工作簿，要求读取数据、完成表格操作并交付修改后的工作簿。"},
                "source": {
                    "repository": REPOSITORY, "commit": COMMIT, "taskPath": f"{ARCHIVE_ROOT}/{prefix}",
                    "files": [{"path": ARCHIVE_PATH, "sha256": ARCHIVE_SHA256}] + [
                        {"path": f"{ARCHIVE_ROOT}/{name}", "sha256": digest(content)}
                        for name, content in [("dataset.json", metadata_bytes), (init_path, init), (gold_path, gold), (prompt_path, prompt)]
                    ],
                    "adaptationChanges": [
                        "从官方 Verified 400 固定快照按显式 ID 选取 12 题；每题包含原版本的一个 init/golden 工作簿对，不冒充完整 912 题三案例设置。",
                        "原始 init 工作簿逐字节复制到 assets/workbook.xlsx，并映射到 input/workbook.xlsx；原始题干保持不变。",
                        "原始 golden 工作簿仅存 private/reference.xlsx；目标单元格缓存值抽取到 private/final.json，不提供给 Agent。",
                        "沿用参考题目的标准 source 字段；来源扩展信息保留在 private/source-metadata.json，结构化目标结果保留在 private/expected.json，final.json 使用完整文本答案。",
                        "要求实际修改并交付 output/result.xlsx，增加过程与本地文件的描述性质证据要求，最终评判遵循 DSHEval LLM 语义评价。",
                        "仅适配数据格式；当前运行器尚无 Question Bundle 加载器，XLSX 证据提取和公式重算后才可完成实际评分；不得等同官方榜单成绩。",
                    ],
                },
                "capabilityLabels": LABELS,
                "task": {"instructions": record["instruction"] + "\n\nThe workbook is at `input/workbook.xlsx`. Apply the requested changes to the workbook and save the completed file as `output/result.xlsx`. Preserve unrelated sheets, cells, and formatting. When a formula is requested, write a working formula. Inspect the saved workbook to verify the result."},
                "environment": {"platform": "portable", "timeoutSeconds": 1200, "dependencies": ["Node.js >=22", "XLSX reading and editing tool", "Excel-compatible formula recalculation engine when output contains formulas"], "inputs": [{"source": "assets/workbook.xlsx", "destination": "input/workbook.xlsx", "sha256": digest(init)}], "setup": {"kind": "none"}, "reset": "fresh-workspace"},
                "final": {"checks": [{"id": "final-workbook", "kind": "llm", "output": "output/result.xlsx", "reference": "private/final.json"}]},
                "evidence": {
                    "process": {"description": "Agent 是否读取实际工作簿、理解跨表或公式约束，并验证修改结果。", "checkpoints": [{"id": "op-1", "description": "读取工作表结构与实际数据，定位任务涉及的区域"}, {"id": "op-2", "description": "按题目约束实施筛选、汇总、匹配、文本处理或公式修改"}, {"id": "op-3", "description": "检查保存后的结果并纠正发现的问题"}]},
                    "local": {"description": "Agent 是否交付可打开的完成工作簿，目标结果正确且无关区域未意外损坏。", "checkpoints": [{"id": "file-1", "description": "创建可打开的 output/result.xlsx"}, {"id": "file-2", "description": "目标区域完成要求的修改，并保留无关工作表与格式"}]},
                },
            }
            expected = {"referenceWorkbook": "private/reference.xlsx", "regions": targets,
                        "upstreamAnswerPosition": record["answer_position"],
                        "upstreamAnswerSheet": record.get("answer_sheet"),
                        "referenceSha256": digest(gold)}
            metadata = {"dataset": "SpreadsheetBench Verified", "upstreamId": task_id, "instructionType": record["instruction_type"], "archiveMemberRoot": ARCHIVE_ROOT, "license": "CC-BY-SA-4.0", "licenseUrl": "https://creativecommons.org/licenses/by-sa/4.0/", "paperUrl": "https://arxiv.org/abs/2406.14991", "testCaseCount": 1, "inputSheetNames": sheets}
            final = {
                "answer": answer_text(expected),
                "rubric": "读取 output/result.xlsx 的实际工作簿内容，与私有 reference.xlsx 及 answer 文本所列 regions 中的目标区域结果作语义比较；相同结构化结果保存在 private/expected.json。需要覆盖全部目标单元格；数值、日期、空值和文本按业务含义比较，不能仅凭文件存在或 Agent 的说明判为正确。题目要求公式时需保留可计算公式，并检查题干明确要求的格式及无关区域。过程和本地层只提供描述性质证据，不绑定固定工具路线。如果无法读取 XLSX 或所需公式结果尚未重算，标记 UNEVALUABLE，不能把缺失缓存当正确空值。此适配不等同官方确定性指标。",
            }
            prepared.append((args.output_root / title, question, final, init, gold, expected, metadata))
        # Validate all source records before creating any bundles.
        for directory, question, final, init, gold, expected, metadata in prepared:
            files = compact_bundle({
                "question.json": json_bytes(question), "private/final.json": json_bytes(final),
                "assets/workbook.xlsx": init, "private/reference.xlsx": gold,
                "private/expected.json": json_bytes(expected), "private/source-metadata.json": json_bytes(metadata),
            })
            for relative, content in files.items():
                write(directory / relative, content)
        write(args.output_root / "LICENSE.md", SPREADSHEET_LICENSE.encode())
    print(json.dumps({"dataset": "spreadsheetbench", "count": len(prepared), "commit": COMMIT, "sourceSha256": ARCHIVE_SHA256}, ensure_ascii=False))


if __name__ == "__main__":
    main()
